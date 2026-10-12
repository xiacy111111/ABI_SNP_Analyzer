# -*- coding: utf-8 -*-
"""Heterozygote-call validation against a manually curated truth table.

Answers SoftwareX reviewer Major 3(b): report heterozygote-specific
sensitivity/precision instead of a single overall concordance rate, and
classify every discordant call.

Usage
-----
    python heterozygote_validation.py \
        --software genotype_matrix.xlsx \
        --truth manual_truth.xlsx \
        --out validation_report.xlsx [--plot confusion_matrix.png]

``--software`` accepts any genotype file ABI-SNP Analyzer writes (native
``genotype_matrix.xlsx``, GENEPOP, Arlequin, hierfstat, GenAlEx, FSTAT,
POPGENE32, STRUCTURE) — format is auto-detected.

``--truth`` is either a matrix (first column = sample, one column per locus)
or a long table with ``Sample`` / ``Locus`` / ``Genotype`` columns.  Genotypes
may be written as ``A/T``, ``AT``, ``T/A`` or as IUPAC ambiguity codes
(W = A/T, R = A/G, Y = C/T, K = G/T, S = C/G, M = A/C).

Definitions (call level, i.e. per sample x locus)
-------------------------------------------------
===================  ========================================================
term                 definition
===================  ========================================================
HET call             the two called alleles differ
HOM call             the two called alleles are identical (or a single allele)
comparable call      non-missing in BOTH the truth table and the software out
sensitivity (het)    TP / (TP + FN) over truth-HET calls
precision (het)      TP / (TP + FP) over software-HET calls
false-het rate       software-HET calls on truth-HOM positions / truth-HOM
===================  ========================================================

Every proportion carries a Wilson score 95% interval, which matters here: with
~880 calls and only a handful of discordances the point estimate is high but
the interval is wide, and the response letter should say so explicitly.
"""

from __future__ import annotations

import argparse
import math
import os
import sys
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Sequence, Tuple

IUPAC = {
    "M": ("A", "C"), "R": ("A", "G"), "W": ("A", "T"),
    "S": ("C", "G"), "Y": ("C", "T"), "K": ("G", "T"),
}
MISSING_TOKENS = {"", "--", "-", "N", "NN", "NA", "N/A", "NAN", ".", "?", "0"}

HET, HOM, MISS = "HET", "HOM", "MISSING"


# ---------------------------------------------------------------------------
# genotype parsing
# ---------------------------------------------------------------------------

def parse_genotype(value) -> Tuple[Optional[str], Optional[str]]:
    """Normalise a genotype cell into two allele letters (or ``(None, None)``)."""
    if value is None:
        return (None, None)
    text = str(value).strip()
    if not text or text.upper() in MISSING_TOKENS:
        return (None, None)
    compact = text.upper().replace("/", "").replace("|", "").replace(" ", "")
    if not compact or compact in MISSING_TOKENS:
        return (None, None)
    if compact in IUPAC:
        return IUPAC[compact]
    if compact in ("A", "C", "G", "T"):
        return (compact, compact)
    if len(compact) == 2 and all(a in "ACGT" for a in compact):
        return (compact[0], compact[1])
    return (None, None)


def zygosity(alleles: Tuple[Optional[str], Optional[str]]) -> str:
    a1, a2 = alleles
    if not a1 or not a2:
        return MISS
    return HET if a1 != a2 else HOM


def genotype_label(alleles: Tuple[Optional[str], Optional[str]]) -> str:
    a1, a2 = alleles
    if not a1 or not a2:
        return "--"
    return "/".join(sorted([a1, a2]))


# ---------------------------------------------------------------------------
# statistics
# ---------------------------------------------------------------------------

def wilson_interval(successes: int, total: int,
                    z: float = 1.959963984540054) -> Tuple[Optional[float],
                                                           Optional[float]]:
    """Wilson score 95% interval — correct for the small discordant counts."""
    if total <= 0:
        return (None, None)
    p = successes / total
    denominator = 1.0 + z * z / total
    centre = (p + z * z / (2 * total)) / denominator
    half = z * math.sqrt(p * (1 - p) / total + z * z / (4 * total * total)) \
        / denominator
    return (max(0.0, centre - half), min(1.0, centre + half))


@dataclass
class ClassMetrics:
    """Sensitivity/precision of one call class (HET or HOM)."""

    class_name: str
    tp: int = 0
    fp: int = 0
    fn: int = 0
    truth_total: int = 0
    predicted_total: int = 0

    @property
    def sensitivity(self) -> Optional[float]:
        return self.tp / self.truth_total if self.truth_total else None

    @property
    def precision(self) -> Optional[float]:
        return self.tp / self.predicted_total if self.predicted_total else None

    @property
    def f1(self) -> Optional[float]:
        if self.sensitivity is None or self.precision is None:
            return None
        if self.sensitivity + self.precision == 0:
            return 0.0
        return 2 * self.sensitivity * self.precision / \
            (self.sensitivity + self.precision)

    @property
    def sensitivity_ci(self):
        return wilson_interval(self.tp, self.truth_total)

    @property
    def precision_ci(self):
        return wilson_interval(self.tp, self.predicted_total)


# ---------------------------------------------------------------------------
# comparison
# ---------------------------------------------------------------------------

@dataclass
class ComparisonResult:
    n_truth_calls: int
    n_software_calls: int
    n_comparable: int
    n_concordant: int
    concordance: Optional[float]
    concordance_ci: Tuple[Optional[float], Optional[float]]
    het: ClassMetrics
    hom: ClassMetrics
    confusion: Dict[str, Dict[str, int]]
    discordant: List[dict] = field(default_factory=list)


def compare_genotypes(truth: Dict[str, Dict[str, Tuple]],
                      software: Dict[str, Dict[str, Tuple]]) -> ComparisonResult:
    """Compare two ``{sample: {locus: (a1, a2)}}`` tables call by call.

    Discordances are classified into:
      * ``missed heterozygote``       truth HET, software HOM
      * ``false heterozygote``        truth HOM, software HET
      * ``homozygote allele mismatch`` both HOM but different alleles
      * ``heterozygote allele mismatch`` both HET but different allele pairs
      * ``software no-call``          truth called, software missing
      * ``truth no-call``             software called, truth missing
    """
    samples = sorted(set(truth) | set(software))
    loci = sorted(set(next(iter(truth.values()), {})) |
                  set(next(iter(software.values()), {})))

    confusion = {truth_cls: {sw_cls: 0 for sw_cls in (HOM, HET, MISS)}
                 for truth_cls in (HOM, HET, MISS)}
    het = ClassMetrics("HET")
    hom = ClassMetrics("HOM")
    discordant: List[dict] = []
    comparable = concordant = 0
    n_truth = n_sw = 0

    CATEGORY = {
        (HET, HOM): "missed heterozygote (software HOM, truth HET)",
        (HOM, HET): "false heterozygote (software HET, truth HOM)",
    }

    for sample in samples:
        for locus in loci:
            truth_call = truth.get(sample, {}).get(locus, (None, None))
            sw_call = software.get(sample, {}).get(locus, (None, None))
            truth_cls = zygosity(truth_call)
            sw_cls = zygosity(sw_call)
            if truth_cls != MISS:
                n_truth += 1
            if sw_cls != MISS:
                n_sw += 1
            confusion[truth_cls][sw_cls] += 1

            if truth_cls == MISS or sw_cls == MISS:
                if truth_cls != sw_cls:
                    discordant.append({
                        "Sample": sample, "Locus": locus,
                        "Truth": genotype_label(truth_call),
                        "Software": genotype_label(sw_call),
                        "Truth_class": truth_cls, "Software_class": sw_cls,
                        "Category": ("software no-call" if sw_cls == MISS
                                     else "truth no-call"),
                    })
                continue

            comparable += 1
            if truth_cls == sw_cls and \
                    genotype_label(truth_call) == genotype_label(sw_call):
                concordant += 1
                continue
            if truth_cls != sw_cls:
                category = CATEGORY[(truth_cls, sw_cls)]
            else:
                category = ("heterozygote allele mismatch" if truth_cls == HET
                            else "homozygote allele mismatch")
            discordant.append({
                "Sample": sample, "Locus": locus,
                "Truth": genotype_label(truth_call),
                "Software": genotype_label(sw_call),
                "Truth_class": truth_cls, "Software_class": sw_cls,
                "Category": category,
            })

    # class tallies come from the confusion matrix (single source of truth):
    #   TP = diagonal, FN = truth-total - TP, FP = predicted-total - TP
    def _fill(metrics: ClassMetrics, cls: str) -> None:
        metrics.truth_total = sum(confusion[cls].values())
        metrics.predicted_total = sum(confusion[t][cls]
                                      for t in (HOM, HET, MISS))
        metrics.tp = confusion[cls][cls]
        metrics.fn = metrics.truth_total - metrics.tp
        metrics.fp = metrics.predicted_total - metrics.tp

    _fill(het, HET)
    _fill(hom, HOM)

    return ComparisonResult(
        n_truth_calls=n_truth,
        n_software_calls=n_sw,
        n_comparable=comparable,
        n_concordant=concordant,
        concordance=concordant / comparable if comparable else None,
        concordance_ci=wilson_interval(concordant, comparable),
        het=het, hom=hom, confusion=confusion, discordant=discordant,
    )


# ---------------------------------------------------------------------------
# input readers
# ---------------------------------------------------------------------------

def _looks_long(frame) -> bool:
    columns = {str(c).strip().lower() for c in frame.columns}
    return {"sample", "locus"} <= columns or {"sample", "snp"} <= columns


def read_truth_table(path: str) -> Dict[str, Dict[str, Tuple]]:
    """Read a manually curated truth table (wide matrix or long format)."""
    import pandas as pd

    if path.lower().endswith((".xlsx", ".xls")):
        frame = pd.read_excel(path, sheet_name=0)
    else:
        frame = pd.read_csv(path, sep=None, engine="python")

    truth: Dict[str, Dict[str, Tuple]] = {}
    if _looks_long(frame):
        sample_col = next(c for c in frame.columns
                          if str(c).strip().lower() == "sample")
        locus_col = next(c for c in frame.columns
                         if str(c).strip().lower() in ("locus", "snp"))
        genotype_col = next((c for c in frame.columns
                             if str(c).strip().lower() in
                             ("genotype", "genotype_call", "call", "truth",
                              "manual")), frame.columns[-1])
        for _, row in frame.iterrows():
            sample = str(row[sample_col]).strip()
            locus = str(row[locus_col]).strip()
            truth.setdefault(sample, {})[locus] = parse_genotype(
                row[genotype_col])
        return truth

    samples = [str(v).strip() for v in frame.iloc[:, 0].tolist()]
    for locus in frame.columns[1:]:
        for sample, value in zip(samples, frame[locus].tolist()):
            truth.setdefault(sample, {})[str(locus).strip()] = \
                parse_genotype(value)
    return truth


def read_software_calls(path: str) -> Dict[str, Dict[str, Tuple]]:
    """Read the software output (any supported genotype format)."""
    try:
        import snp_format_io as snpio
    except ImportError:
        snpio = None
    if snpio is not None:
        data = snpio.read_genotype_file(path)
        calls = data.to_matrix_dict()          # {locus: {sample: "AT"/"--"}}
        return {
            sample: {locus: parse_genotype(calls[locus][sample])
                     for locus in calls}
            for sample in data.samples
        }
    import pandas as pd
    frame = pd.read_excel(path, sheet_name="Genotype_Matrix")
    result: Dict[str, Dict[str, Tuple]] = {}
    samples = [str(v).strip() for v in frame.iloc[:, 0].tolist()]
    for locus in frame.columns[1:]:
        for sample, value in zip(samples, frame[locus].tolist()):
            result.setdefault(sample, {})[str(locus).strip()] = \
                parse_genotype(value)
    return result


# ---------------------------------------------------------------------------
# report
# ---------------------------------------------------------------------------

def write_report(path: str, result: ComparisonResult,
                 truth_path: str, software_path: str) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Summary"
    sheet.column_dimensions["A"].width = 52
    sheet.column_dimensions["B"].width = 24
    rows = [
        ("Genotype-call validation", ""),
        ("Truth table", os.path.basename(truth_path)),
        ("Software output", os.path.basename(software_path)),
        ("Truth calls (non-missing)", result.n_truth_calls),
        ("Software calls (non-missing)", result.n_software_calls),
        ("Comparable calls (both non-missing)", result.n_comparable),
        ("Concordant calls", result.n_concordant),
        ("Overall concordance", _pct(result.concordance)),
        ("Overall concordance 95% CI (Wilson)",
         _ci(result.concordance_ci)),
        ("", ""),
        ("Heterozygote sensitivity (TP/(TP+FN))",
         _pct(result.het.sensitivity)),
        ("Heterozygote sensitivity 95% CI", _ci(result.het.sensitivity_ci)),
        ("Heterozygote precision (TP/(TP+FP))",
         _pct(result.het.precision)),
        ("Heterozygote precision 95% CI", _ci(result.het.precision_ci)),
        ("Heterozygote F1", _pct(result.het.f1)),
        ("False-heterozygote rate on truth-HOM positions",
         _fraction(result.confusion[HOM][HET], result.hom.truth_total)),
        ("", ""),
        ("Homozygote sensitivity", _pct(result.hom.sensitivity)),
        ("Homozygote precision", _pct(result.hom.precision)),
        ("", ""),
        ("Discordant calls", len(result.discordant)),
    ]
    for row in rows:
        sheet.append(list(row))
    sheet["A1"].font = Font(bold=True)
    for row in sheet.iter_rows(min_row=2):
        if row[1].value in (None, ""):
            row[0].font = Font(bold=True)

    sheet = workbook.create_sheet("Confusion_Matrix")
    sheet.append(["Truth \\ Software", "HOM", "HET", "MISSING", "Total"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for truth_cls in (HOM, HET, MISS):
        row = [truth_cls]
        for sw_cls in (HOM, HET, MISS):
            row.append(result.confusion[truth_cls][sw_cls])
        row.append(sum(result.confusion[truth_cls].values()))
        sheet.append(row)
    sheet.append(["Total"] + [sum(result.confusion[t][s] for t in
                                  (HOM, HET, MISS))
                              for s in (HOM, HET, MISS)])

    sheet = workbook.create_sheet("Metrics")
    sheet.append(["Class", "TP", "FP", "FN", "Truth total",
                  "Predicted total", "Sensitivity", "Sensitivity 95% CI",
                  "Precision", "Precision 95% CI", "F1"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for metrics in (result.het, result.hom):
        sheet.append([
            metrics.class_name, metrics.tp, metrics.fp, metrics.fn,
            metrics.truth_total, metrics.predicted_total,
            _pct(metrics.sensitivity), _ci(metrics.sensitivity_ci),
            _pct(metrics.precision), _ci(metrics.precision_ci),
            _pct(metrics.f1)])

    sheet = workbook.create_sheet("Discordant_Calls")
    headers = ["Sample", "Locus", "Truth", "Software", "Truth_class",
               "Software_class", "Category"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in result.discordant:
        sheet.append([row[h] for h in headers])
    _autosize(sheet)

    workbook.save(path)
    return path


def _pct(value: Optional[float]) -> str:
    return "" if value is None else f"{value:.4f}"


def _ci(interval) -> str:
    if interval is None or interval[0] is None:
        return ""
    return f"[{interval[0]:.4f}, {interval[1]:.4f}]"


def _fraction(numerator: int, denominator: int) -> str:
    if not denominator:
        return ""
    return f"{numerator / denominator:.4f} ({numerator}/{denominator})"


def _autosize(sheet) -> None:
    for column_cells in sheet.columns:
        length = max(len(str(cell.value or "")) for cell in column_cells)
        letter = column_cells[0].column_letter
        sheet.column_dimensions[letter].width = min(max(length + 2, 10), 48)


def plot_confusion(result: ComparisonResult, path: str) -> str:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import numpy as np

    classes = (HOM, HET, MISS)
    matrix = np.array([[result.confusion[t][s] for s in classes]
                       for t in classes], dtype=float)
    figure, axis = plt.subplots(figsize=(5.2, 4.4))
    image = axis.imshow(matrix, cmap="Blues")
    axis.set_xticks(range(3), ["Software\nHOM", "Software\nHET",
                               "Software\nmissing"])
    axis.set_yticks(range(3), ["Truth\nHOM", "Truth\nHET", "Truth\nmissing"])
    for i in range(3):
        for j in range(3):
            axis.text(j, i, int(matrix[i, j]), ha="center", va="center",
                      color="white" if matrix[i, j] > matrix.max() / 2
                      else "black")
    axis.set_title("Genotype-call confusion matrix")
    figure.colorbar(image, ax=axis, shrink=0.8)
    figure.tight_layout()
    figure.savefig(path, dpi=300)
    plt.close(figure)
    return path


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(
        description="Validate SNP calls against a manually curated truth table.")
    parser.add_argument("--software", required=True,
                        help="genotype file produced by ABI-SNP Analyzer")
    parser.add_argument("--truth", required=True,
                        help="manually curated truth table (.xlsx/.csv)")
    parser.add_argument("--out", default="validation_report.xlsx",
                        help="output Excel report")
    parser.add_argument("--plot", default=None,
                        help="optional confusion-matrix PNG")
    args = parser.parse_args(argv)

    truth = read_truth_table(args.truth)
    software = read_software_calls(args.software)
    result = compare_genotypes(truth, software)

    print(f"Comparable calls      : {result.n_comparable}")
    print(f"Concordance           : "
          f"{_pct(result.concordance)} {_ci(result.concordance_ci)}")
    print(f"HET sensitivity       : {_pct(result.het.sensitivity)} "
          f"({_result_counts(result.het)}) {_ci(result.het.sensitivity_ci)}")
    print(f"HET precision         : {_pct(result.het.precision)} "
          f"({_result_counts(result.het)}) {_ci(result.het.precision_ci)}")
    print(f"Discordant calls      : {len(result.discordant)}")
    for row in result.discordant:
        print(f"  {row['Sample']} @ {row['Locus']}: "
              f"truth {row['Truth']} vs software {row['Software']} "
              f"-> {row['Category']}")

    write_report(args.out, result, args.truth, args.software)
    print(f"Report written        : {args.out}")
    if args.plot:
        plot_confusion(result, args.plot)
        print(f"Figure written        : {args.plot}")
    return 0


def _result_counts(metrics: ClassMetrics) -> str:
    return f"TP={metrics.tp}, FP={metrics.fp}, FN={metrics.fn}"


if __name__ == "__main__":
    sys.exit(main())
